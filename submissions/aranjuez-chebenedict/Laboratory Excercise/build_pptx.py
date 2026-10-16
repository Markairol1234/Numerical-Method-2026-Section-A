"""
Build the PowerPoint presentation for Series 1.

Slides:
  1. Title
  2. The constant e (overview)
  3. Exercise 1 - concept + answer
  4. Exercise 2 - concept + answer
  5. Exercise 3 - concept + answer
  6. Comparison & conclusion

Each exercise slide includes the generated PNG figure and a text explanation.
Run after the three exercise scripts (they create the *_output.png files).
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK = RGBColor(0x0F, 0x17, 0x2A)
ACCENT = RGBColor(0x38, 0xBD, 0xF8)
TEXT = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BASE = r"d:\che school 2627\lab04"


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK
    bg.line.fill.background()
    bg.shadow.inherit = False

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.4),
                                  Inches(8.4), Inches(1.6))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = ACCENT

    tb2 = slide.shapes.add_textbox(Inches(1.0), Inches(4.0),
                                   Inches(8.0), Inches(1.2))
    p2 = tb2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = subtitle
    r2.font.size = Pt(20); r2.font.color.rgb = TEXT
    return slide


def add_content_slide(prs, heading, bullets, image_path=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.1))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK
    bg.line.fill.background()
    bg.shadow.inherit = False

    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15),
                                  Inches(12.5), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = heading
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = ACCENT

    top = Inches(1.3)
    if image_path:
        box = slide.shapes.add_textbox(Inches(0.4), top,
                                       Inches(5.6), Inches(5.6))
        slide.shapes.add_picture(image_path, Inches(6.2), Inches(1.5),
                                 width=Inches(6.9))
    else:
        box = slide.shapes.add_textbox(Inches(0.7), top,
                                       Inches(11.9), Inches(5.5))

    tf = box.text_frame
    tf.word_wrap = True
    for i, (level, text) in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.level = level
        run = para.add_run(); run.text = text
        run.font.size = Pt(17 if level == 0 else 14)
        run.font.color.rgb = WHITE if level == 0 else TEXT
        if level == 0:
            run.font.bold = True
        para.space_after = Pt(5)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Series 1 - Laboratory Exercise",
        "Computing Euler's number e three ways: compound interest, "
        "the difference quotient, and the exponential series")

    add_content_slide(prs, "The Constant e  (overview)", [
        (0, "e = 2.718281828459045..."),
        (1, "Irrational, transcendental - the base of natural logarithms."),
        (1, "Appears in growth, calculus, probability and complex analysis."),
        (0, "Three ways to reach it in this lab:"),
        (1, "1. Compound-interest limit  (1 + 1/n)^n  ->  e"),
        (1, "2. Derivative  d/dx[a^x] at x=0 = ln(a)  via a difference quotient"),
        (1, "3. Exponential Taylor series  e^x = sum x^n / n!"),
        (0, "Goal: compare convergence speed and the float-precision limit."),
    ])

    add_content_slide(prs, "Exercise 1 - (1 + 1/n)^n -> e", [
        (0, "Concept"),
        (1, "Compound a 100% rate n times per year: factor = (1 + 1/n)^n."),
        (1, "As n grows (yearly -> every nanosecond) the factor -> e."),
        (0, "Answer / Result"),
        (1, "yearly n=1 -> 2.0000 ; daily n=365 -> 2.7146"),
        (1, "by every second it is stable to 7 decimals (2.718282)."),
        (1, "Error falls like e/(2n): doubling n roughly halves the error."),
        (0, "Numerical trick"),
        (1, "Use exp(n*log1p(1/n)); the naive form rounds to 1.0 for n >= 1e16."),
    ], image_path=rf"{BASE}\exercise1_output.png")

    add_content_slide(prs, "Exercise 2 - (a^h - 1)/h -> ln(a)", [
        (0, "Concept"),
        (1, "This is the derivative of a^x at x = 0, i.e. ln(a)."),
        (1, "The forward difference quotient is a finite-step estimate of it."),
        (0, "Answer / Result"),
        (1, "a=2 settles at 0.6931  (ln 2)"),
        (1, "a=e settles at 1.0000  (ln e = 1, the 'natural' base)"),
        (1, "a=3 settles at 1.0986  (ln 3)"),
        (1, "Truncation error is O(h): each 10x smaller h adds ~1 digit."),
        (0, "Numerical trick"),
        (1, "Use expm1(h*ln(a))/h to avoid catastrophic cancellation."),
    ], image_path=rf"{BASE}\exercise2_output.png")

    add_content_slide(prs, "Exercise 3 - e^x = sum x^n / n!", [
        (0, "Concept"),
        (1, "Maclaurin series of exp(x); for x=1 it gives the series for e."),
        (1, "Terms 1/n! vanish super fast, so convergence is super-linear."),
        (0, "Answer / Result"),
        (1, "N=10 -> ~8 correct digits ; N=15 -> ~13 ; N=20 -> full double precision."),
        (1, "Beyond N~20 the error is frozen at the machine floor 4.44e-16."),
        (1, "Summing to N=10,000 gives the same result as N=20."),
        (0, "Why it flattens"),
        (1, "Once 1/n! < machine epsilon (~2.2e-16) extra terms change nothing."),
    ], image_path=rf"{BASE}\exercise3_output.png")

    add_content_slide(prs, "Comparison & Conclusion", [
        (0, "Three routes to e, three convergence speeds:"),
        (1, "Ex.1  (1+1/n)^n    : O(1/n)       - needs n ~ 1e15 for 15 digits (slowest)"),
        (1, "Ex.2  (a^h-1)/h    : O(h) / decade- needs h ~ 1e-15; hits cancellation limit"),
        (1, "Ex.3  sum x^n/n!   : super-linear - only ~20 terms for full precision (fastest)"),
        (0, "Conclusion"),
        (1, "The Taylor series is the most efficient; it is what computers use for exp(x)."),
        (1, "The limits (Ex.1, Ex.2) converge too slowly for practical computation."),
        (1, "Numerically-stable formulas (log1p, expm1) are essential near the precision limit."),
    ])

    out = rf"{BASE}\Series1_Presentation.pptx"
    prs.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
