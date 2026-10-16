# ============================================================
# cube_html.py
# Builds a 6m x 6m x 6m cube structural model and writes:
#   1) cube_model.xlsx   -> "Nodes" + "Member Incidences" sheets
#   2) cube_model.html   -> self-contained HTML report with the
#                           3D wireframe (base64) and tables.
#
# Axis convention (RISA-3D / STAAD.Pro compatible):
#    Global X, Z = lateral (horizontal)
#    Global Y    = vertical (up)
# ============================================================

import os
import base64
import datetime as dt
import io

import matplotlib
matplotlib.use("Agg")             # no pop-up — save to buffer
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# ------------------------------------------------------------
# 0. Paths
# ------------------------------------------------------------
OUT_DIR    = r"D:\che school 2627\lab02"
EXCEL_PATH = os.path.join(OUT_DIR, "cube_model.xlsx")
HTML_PATH  = os.path.join(OUT_DIR, "cube_model.html")
os.makedirs(OUT_DIR, exist_ok=True)

# ------------------------------------------------------------
# 1. Cube geometry
# ------------------------------------------------------------
L = 6.0

nodes = {
    1: (0.0, 0.0, 0.0),
    2: (L,   0.0, 0.0),
    3: (L,   0.0, L),
    4: (0.0, 0.0, L),
    5: (0.0, L,   0.0),
    6: (L,   L,   0.0),
    7: (L,   L,   L),
    8: (0.0, L,   L),
}

member_pairs = [
    (1, 2), (2, 3), (3, 4), (4, 1),   # bottom face
    (5, 6), (6, 7), (7, 8), (8, 5),   # top face
    (1, 5), (2, 6), (3, 7), (4, 8),   # verticals
]

def classify_member(i, j):
    xi, yi, zi = nodes[i]
    xj, yj, zj = nodes[j]
    if xi == xj and zi == zj and yi != yj:
        return "Column"
    return "Beam"

members = [(i, j, classify_member(i, j)) for i, j in member_pairs]

# ------------------------------------------------------------
# 2. Excel workbook — Nodes and Member Incidences
# ------------------------------------------------------------
wb = openpyxl.Workbook()

header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
body_font   = Font(name="Arial")
center      = Alignment(horizontal="center")

def style_header(ws, headers):
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center

def style_body(ws, ncols):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
        for cell in row:
            cell.font = body_font
            cell.alignment = center

# Sheet 1: Nodes
ws_nodes = wb.active
ws_nodes.title = "Nodes"
style_header(ws_nodes, ["Node", "X (m)", "Y (m)", "Z (m)"])
for n, (x, y, z) in nodes.items():
    ws_nodes.append([n, x, y, z])
style_body(ws_nodes, 4)
for col_letter, width in zip("ABCD", [10, 12, 12, 12]):
    ws_nodes.column_dimensions[col_letter].width = width

# Sheet 2: Member Incidences
ws_mem = wb.create_sheet("Member Incidences")
style_header(ws_mem, ["Member", "i (Start Node)", "j (End Node)", "Type"])
for idx, (i, j, mtype) in enumerate(members, start=1):
    ws_mem.append([idx, i, j, mtype])
style_body(ws_mem, 4)
for col_letter, width in zip("ABCD", [10, 18, 18, 12]):
    ws_mem.column_dimensions[col_letter].width = width

wb.save(EXCEL_PATH)
print(f"Saved Excel: {EXCEL_PATH}")

# ------------------------------------------------------------
# 3. 3D plot -> PNG buffer -> base64 <img>
#    NOTE: matplotlib's 3D axes take (x, y, z) in ITS OWN order
#    where z is drawn as "up". To make Y the visual vertical
#    (solver convention), we pass model (X, Z, Y) as (x, y, z).
# ------------------------------------------------------------
fig = plt.figure(figsize=(9, 9))
ax  = fig.add_subplot(111, projection="3d")

for idx, (i, j, mtype) in enumerate(members, start=1):
    x_vals = [nodes[i][0], nodes[j][0]]
    z_vals = [nodes[i][2], nodes[j][2]]
    y_vals = [nodes[i][1], nodes[j][1]]
    color = "darkorange" if mtype == "Column" else "steelblue"
    ax.plot(x_vals, z_vals, y_vals, color=color, linewidth=2)

# nodes
xs = [c[0] for c in nodes.values()]
zs = [c[2] for c in nodes.values()]
ys = [c[1] for c in nodes.values()]
ax.scatter(xs, zs, ys, color="red", s=60, depthshade=True)

for n, (x, y, z) in nodes.items():
    ax.text(x, z, y, f"  {n}", color="black", fontsize=10, fontweight="bold")

# global axis triad at origin
axis_len = L * 0.4
ax.quiver(0, 0, 0, axis_len, 0, 0, color="black", linewidth=2, arrow_length_ratio=0.12)
ax.quiver(0, 0, 0, 0, axis_len, 0, color="black", linewidth=2, arrow_length_ratio=0.12)
ax.quiver(0, 0, 0, 0, 0, axis_len, color="black", linewidth=2, arrow_length_ratio=0.12)
ax.text(axis_len * 1.05, 0, 0, "Global X", fontsize=9, fontweight="bold")
ax.text(0, axis_len * 1.05, 0, "Global Z", fontsize=9, fontweight="bold")
ax.text(0, 0, axis_len * 1.05, "Global Y", fontsize=9, fontweight="bold")

ax.set_xlabel("X (m) - Lateral",  fontsize=11, labelpad=10)
ax.set_ylabel("Z (m) - Lateral",  fontsize=11, labelpad=10)
ax.set_zlabel("Y (m) - Vertical", fontsize=11, labelpad=10)

ax.set_title("6m x 6m x 6m Cube Model\n"
             "(X, Z = Lateral  |  Y = Global Vertical)",
             fontsize=12, fontweight="bold")

ax.set_xlim(-1, L + 1)
ax.set_ylim(-1, L + 1)
ax.set_zlim(0, L + 1)
ax.set_box_aspect([1, 1, 1])
ax.view_init(elev=20, azim=-60)

legend_elems = [
    Line2D([0], [0], color="steelblue",  lw=2, label="Beam"),
    Line2D([0], [0], color="darkorange", lw=2, label="Column"),
]
ax.legend(handles=legend_elems, loc="upper left", fontsize=9)

plt.tight_layout()
buf = io.BytesIO()
fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
plt.close(fig)
buf.seek(0)
IMG_B64 = base64.b64encode(buf.read()).decode("ascii")

# ------------------------------------------------------------
# 4. HTML report
# ------------------------------------------------------------
generated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")

# --- node rows ---
node_rows = ""
for n, (x, y, z) in nodes.items():
    node_rows += f"<tr><td>{n}</td><td>{x:.1f}</td><td>{y:.1f}</td><td>{z:.1f}</td></tr>"

# --- member rows ---
member_rows = ""
for idx, (i, j, mtype) in enumerate(members, start=1):
    cls = "col-beam" if mtype == "Beam" else "col-column"
    member_rows += (f"<tr><td>{idx}</td><td>{i}</td><td>{j}</td>"
                    f"<td class='{cls}'>{mtype}</td></tr>")

CSS = """
:root {
  --ink:#0b2531; --muted:#5b7180; --line:#d7e3e6;
  --teal:#0e7c86; --teal-dark:#0b4f6c; --amber:#ffb238;
}
* { box-sizing: border-box; }
body {
  margin:0; font-family:'Segoe UI', Arial, sans-serif; color:var(--ink);
  background:#f2fafa; line-height:1.55;
}
header.hero {
  background:linear-gradient(135deg,#0e7c86 0%, #0b4f6c 100%);
  color:#fff; padding:32px 28px 60px 28px; position:relative;
}
header.hero .kicker {
  display:inline-block; background:rgba(255,255,255,0.15); padding:4px 10px;
  border-radius:20px; font-size:11px; letter-spacing:0.08em; text-transform:uppercase;
}
header.hero h1 { margin:12px 0 6px 0; font-size:32px; }
header.hero .sub { opacity:0.85; font-size:14px; }
.wave { position:absolute; left:0; right:0; bottom:-2px; height:60px; }
.wave svg { width:100%; height:100%; display:block; }
main { max-width:1100px; margin:0 auto; padding:24px 20px 40px; }
section.panel {
  background:#fff; border-radius:16px; padding:22px 24px; margin-bottom:22px;
  box-shadow:0 6px 24px rgba(11,79,108,0.08);
}
h2 { margin:0 0 12px 0; color:var(--teal-dark); font-size:20px; }
h3.label {
  font-size:11px; letter-spacing:0.08em; text-transform:uppercase;
  color:var(--muted); margin:18px 0 8px 0; font-weight:700;
}
.chart { max-width:100%; height:auto; border-radius:10px; display:block; margin:0 auto; }
table { border-collapse:collapse; width:100%; font-size:13px; margin:10px 0; }
th, td { border-bottom:1px solid var(--line); padding:9px 10px; text-align:left; }
th {
  background:#effaf9; color:var(--teal-dark); text-transform:uppercase;
  font-size:10.5px; letter-spacing:0.05em;
}
tr:last-child td { border-bottom:none; }
.col-beam   { color:#2a7de1; font-weight:700; }
.col-column { color:#c25e00; font-weight:700; }
.stat-grid {
  display:grid; grid-template-columns:repeat(auto-fit,minmax(180px,1fr));
  gap:14px; margin:14px 0;
}
.stat-card {
  border-radius:14px; padding:12px 14px; background:#f7fcfc;
  border-left:5px solid var(--teal); box-shadow:0 4px 14px rgba(11,79,108,0.06);
}
.stat-card .label {
  font-size:10.5px; text-transform:uppercase; color:var(--muted);
  font-weight:700; letter-spacing:0.06em;
}
.stat-card .value { font-size:20px; font-weight:800; color:var(--ink); margin-top:4px; }
.axes-note {
  background:#effaf9; border-left:4px solid var(--teal);
  padding:10px 14px; border-radius:8px; font-size:13px; margin:10px 0;
}
footer { text-align:center; color:var(--muted); font-size:11.5px; padding:14px; }
"""

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>6m Cube Structural Model</title>
<style>
{css}
</style>
</head>
<body>

<header class="hero">
  <span class="kicker">Structural Model &middot; 6m Cube</span>
  <h1>6 m &times; 6 m &times; 6 m Cube Frame</h1>
  <div class="sub">
    Node coordinates and member incidences &middot;
    Axis convention: X, Z = lateral &middot; Y = vertical &middot;
    generated {generated}
  </div>
  <div class="wave">
    <svg viewBox="0 0 1200 60" preserveAspectRatio="none">
      <path d="M0,32 C150,60 350,0 600,28 C850,56 1050,8 1200,30 L1200,60 L0,60 Z"
            fill="#f2fafa"/>
    </svg>
  </div>
</header>

<main>

  <section class="panel">
    <h2>Model Overview</h2>
    <div class="stat-grid">
      <div class="stat-card"><div class="label">Side Length</div><div class="value">{L:.1f} m</div></div>
      <div class="stat-card"><div class="label">Nodes</div><div class="value">{n_nodes}</div></div>
      <div class="stat-card"><div class="label">Members</div><div class="value">{n_members}</div></div>
      <div class="stat-card"><div class="label">Beams / Columns</div><div class="value">{n_beam} / {n_col}</div></div>
    </div>
    <div class="axes-note">
      <strong>Axis convention (RISA-3D / STAAD.Pro compatible):</strong>
      Global <strong>X</strong> and <strong>Z</strong> are lateral (horizontal),
      Global <strong>Y</strong> is vertical up. Node 1 sits at the origin (0, 0, 0).
    </div>
  </section>

  <section class="panel">
    <h2>3D Wireframe</h2>
    <p>Blue = Beam &middot; Orange = Column &middot; Black arrows = global X, Y, Z axes at the origin.</p>
    <img class="chart" src="data:image/png;base64,{img_b64}" alt="6m cube wireframe">
  </section>

  <section class="panel">
    <h2>Nodes</h2>
    <table>
      <tr><th>Node</th><th>X (m)</th><th>Y (m)</th><th>Z (m)</th></tr>
      {node_rows}
    </table>
  </section>

  <section class="panel">
    <h2>Member Incidences</h2>
    <p>i = start node &middot; j = end node.</p>
    <table>
      <tr><th>Member</th><th>i (Start Node)</th><th>j (End Node)</th><th>Type</th></tr>
      {member_rows}
    </table>
  </section>

</main>

<footer>
  Generated by <code>cube_html.py</code> &mdash;
  every value above was computed in Python; the page contains no live computation.
</footer>

</body>
</html>
"""

html = HTML_TEMPLATE.format(
    css=CSS,
    generated=generated,
    L=L,
    n_nodes=len(nodes),
    n_members=len(members),
    n_beam=sum(1 for _, _, t in members if t == "Beam"),
    n_col=sum(1 for _, _, t in members if t == "Column"),
    img_b64=IMG_B64,
    node_rows=node_rows,
    member_rows=member_rows,
)

with open(HTML_PATH, "w", encoding="utf-8") as f:
    f.write(html)

print(f"Saved HTML:  {HTML_PATH}")