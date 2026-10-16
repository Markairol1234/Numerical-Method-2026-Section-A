# ============================================================
# Rev2.py  —  Rev. 2
# 6m x 6m x 6m Cube Structural Model Generator
#
# PURPOSE
#   Build a 6 m cube space-frame model with pinned supports, beta
#   angles, RISA/STAAD-compatible local axes (L1 axial, L2/L3 cross-
#   section), nodal DOF numbering and member end releases; then emit
#   an Excel workbook, a self-contained HTML report and a 3D matplotlib
#   render. Rev2 adds unit-system switching (metric default / imperial),
#   an A36 Steel material and a W 100x100 (= W 4x13) section, all shown
#   on the render and in the outputs.
#
# AXIS CONVENTION (RISA / STAAD compatible)
#   X, Z = lateral;  Y = vertical (up).
#   Local-1 = axial, Local-2 / Local-3 = cross-section axes.
#
# UNIT SYSTEM
#   UNIT_SYSTEM = "metric"   -> length m, force kN, stress MPa,
#                               section mm, label "W 100x100".
#   UNIT_SYSTEM = "imperial" -> length ft, force kip, stress ksi,
#                               section in,  label "W 4x13".
#   All physics is stored in SI internally; only display/input converts.
#
# MATERIAL   A36 Steel (E, Fy, Fu, rho, nu, G, alpha).
# SECTION    W 100x100 metric (== imperial W 4x13), AISC DB v16.0.
# ============================================================

import os
import base64
import datetime as dt
import io

import numpy as np
import matplotlib
# Use Agg if running headless (no display / CI); TkAgg enables the pop-up.
matplotlib.use("TkAgg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

# ============================================================
# 1. CONFIG — units, material, section, geometry
# ============================================================
UNIT_SYSTEM = "metric"          # "metric" (default) | "imperial"

# ---- 1a. Unit conversion constants (NIST SP 811 exact defs) -------
# Declared ONCE here. No inline conversion factors anywhere else.
UNITS = {
    # base exact definitions
    "IN_TO_MM_EXACT": 25.4,
    "LBF_TO_N_EXACT": 4.4482216152605,
    "GRAVITY_M_S2": 9.80665,
    # length
    "M_TO_MM": 1000.0,
    "M2_TO_MM2": 1000.0 ** 2,                 # 1.0e6
    "M4_TO_MM4": 1000.0 ** 4,                 # 1.0e12
    "MM_TO_IN": 0.03937007874015748,
    "IN_TO_MM": 25.4,
    "M_TO_FT": 3.280839895013123,
    "FT_TO_M": 0.3048,
    # area
    "IN2_TO_MM2": 645.16,
    "MM2_TO_IN2": 0.0015500031000062,
    # moment of inertia
    "IN4_TO_MM4": 416231.4256,
    "MM4_TO_IN4": 2.4025096100288302e-06,
    # section modulus
    "IN3_TO_MM3": 16387.064,
    "MM3_TO_IN3": 6.102374409473229e-05,
    # force
    "LBF_TO_N": 4.4482216152605,
    "N_TO_LBF": 0.2248089430997105,
    "KIP_TO_N": 4448.2216152605,
    "KIP_TO_KN": 4.4482216152605,
    "N_TO_KIP": 2.2480894309971046e-04,
    "N_PER_KN": 1000.0,                       # SI prefix (exact)
    # stress / modulus
    "PSI_TO_PA": 6894.757293168361,
    "PA_TO_PSI": 0.00014503773773020923,
    "KSI_TO_MPA": 6.894757293168361,
    "MPA_TO_KSI": 0.14503773773020923,
    "PA_PER_MPA": 1.0e6,                      # SI prefix (exact)
    # unit weight (display only)
    "PCF_TO_KN_M3": 0.15708746384624625,
    "KFT3_TO_KN_M3": 157.08746384624624,
    # mass (display only) — avoirdupois pound, exact
    "LB_TO_KG": 0.45359237,
    "KG_TO_LB": 2.20462262184878,
}

# ---- 1b. Material — A36 Steel (SI internal: Pa, kg/m^3) -----------
MATERIAL = {
    "name": "A36 Steel",
    "E": 199948e6,        # Pa   (199.948 GPa) — RISA Materials Library (Metric)
    "Fy": 248.2e6,        # Pa   (248.2 MPa)   yield strength
    "Fu": 399.9e6,        # Pa   (399.9 MPa)   ultimate strength
    "rho": 7848.76,       # kg/m^3             mass density
    "nu": 0.30,           # —                  Poisson's ratio
    "G": 76904e6,         # Pa   (76.904 GPa)  shear modulus
    "alpha": 11.7e-6,     # 1/°C               thermal coefficient
}

# ---- 1c. Section — W 100x100 metric (== W 4x13), SI internal -------
SECTION = {
    "name": "W 100x100",
    "A": 2.470e-3,        # m^2   cross-sectional area
    "Ix": 4.70e-6,        # m^4   strong-axis moment of inertia
    "Iy": 1.61e-6,        # m^4   weak-axis moment of inertia
    "J": 6.29e-8,         # m^4   torsion constant
    "d": 0.106,           # m     depth
    "bf": 0.103,          # m     flange width
    "tf": 0.00876,        # m     flange thickness
    "tw": 0.00711,        # m     web thickness
}

# Optional per-member section override, e.g. {7: SECTION_ALT}. Empty by
# default; every member uses SECTION. Keys are 1-based member indices.
SECTION_OVERRIDES = {}

# ---- 1d. Geometry — 6 m cube --------------------------------------
L = 6.0                   # cube side length (m)

nodes = {
    1: (0.0, 0.0, 0.0),
    2: (L,   0.0, 0.0),
    3: (L,   0.0, L),
    4: (0.0, 0.0, L),
    5: (0.0, L,   0.0),
    6: (L,   L,   0.0),
    7: (L,   L,   L),
    8: (0.0, L,   L),
}

BASE_NODES = (1, 2, 3, 4)  # pinned supports

member_pairs = [
    (1, 2), (2, 3), (3, 4), (4, 1),
    (5, 6), (6, 7), (7, 8), (8, 5),
    (1, 5), (2, 6), (3, 7), (4, 8),
]

DOF_NAMES = ["Ux", "Uy", "Uz", "Rx", "Ry", "Rz"]
LOCAL_DOF_NAMES = ["Fx", "Fy", "Fz", "Mx", "My", "Mz"]

# ---- 1e. Output paths ---------------------------------------------
OUT_DIR = r"D:\che school 2627\lab02\Rev2"
EXCEL_PATH = os.path.join(OUT_DIR, "cube_model.xlsx")
HTML_PATH = os.path.join(OUT_DIR, "cube_model.html")
os.makedirs(OUT_DIR, exist_ok=True)


# ============================================================
# 2. UNITS — display helpers (SI internal -> display, and inverse)
# ============================================================
def get_unit_labels():
    """Return display unit labels for the active UNIT_SYSTEM."""
    if UNIT_SYSTEM == "imperial":
        return {"length": "ft", "force": "kip", "stress": "ksi",
                "section_dim": "in", "area": "in²", "inertia": "in⁴",
                "mass": "lb", "unit_weight": "k/ft³"}
    return {"length": "m", "force": "kN", "stress": "MPa",
            "section_dim": "mm", "area": "mm²", "inertia": "mm⁴",
            "mass": "kg", "unit_weight": "kN/m³"}


UNIT_LABELS = get_unit_labels()

# Factor converting internal metres -> plotting/display length units.
LEN_TO_DISPLAY = UNITS["M_TO_FT"] if UNIT_SYSTEM == "imperial" else 1.0


def to_display_length(m):
    """Length: internal metres -> display unit (m or ft)."""
    if UNIT_SYSTEM == "imperial":
        return m * UNITS["M_TO_FT"]
    return m


def from_display_length(val):
    """Length: display unit (m or ft) -> internal metres."""
    if UNIT_SYSTEM == "imperial":
        return val * UNITS["FT_TO_M"]
    return val


def to_display_force(n):
    """Force: internal newtons -> display unit (kN or kip)."""
    if UNIT_SYSTEM == "imperial":
        return n * UNITS["N_TO_KIP"]
    return n / UNITS["N_PER_KN"]


def from_display_force(val):
    """Force: display unit (kN or kip) -> internal newtons."""
    if UNIT_SYSTEM == "imperial":
        return val / UNITS["N_TO_KIP"]
    return val * UNITS["N_PER_KN"]


def to_display_stress(pa):
    """Stress/modulus: internal pascals -> display unit (MPa or ksi)."""
    if UNIT_SYSTEM == "imperial":
        return (pa / UNITS["PA_PER_MPA"]) * UNITS["MPA_TO_KSI"]
    return pa / UNITS["PA_PER_MPA"]


def from_display_stress(val):
    """Stress/modulus: display unit (MPa or ksi) -> internal pascals."""
    if UNIT_SYSTEM == "imperial":
        return (val / UNITS["MPA_TO_KSI"]) * UNITS["PA_PER_MPA"]
    return val * UNITS["PA_PER_MPA"]


def to_display_section_dim(m):
    """Section dimension: internal metres -> display unit (mm or in)."""
    if UNIT_SYSTEM == "imperial":
        return m * UNITS["M_TO_MM"] * UNITS["MM_TO_IN"]
    return m * UNITS["M_TO_MM"]


def to_display_section_area(m2):
    """Section area: internal m^2 -> display unit (mm² or in²)."""
    if UNIT_SYSTEM == "imperial":
        return m2 * UNITS["M2_TO_MM2"] * UNITS["MM2_TO_IN2"]
    return m2 * UNITS["M2_TO_MM2"]


def to_display_section_inertia(m4):
    """Section inertia: internal m^4 -> display unit (mm⁴ or in⁴)."""
    if UNIT_SYSTEM == "imperial":
        return m4 * UNITS["M4_TO_MM4"] * UNITS["MM4_TO_IN4"]
    return m4 * UNITS["M4_TO_MM4"]


def to_display_mass(kg):
    """Mass: internal kg -> display unit (kg or lb)."""
    if UNIT_SYSTEM == "imperial":
        return kg * UNITS["KG_TO_LB"]
    return kg


def to_plot(x):
    """Internal metres -> plotting/display length units."""
    return x * LEN_TO_DISPLAY


def section_label():
    """Section name for the active unit system."""
    return "W 4x13" if UNIT_SYSTEM == "imperial" else "W 100x100"


def get_section(idx):
    """Section dict for member idx (honours SECTION_OVERRIDES)."""
    return SECTION_OVERRIDES.get(idx, SECTION)


def fmt_value(v):
    """Format a display value (int if whole, else 4 significant figures)."""
    if isinstance(v, str):
        return v
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (int, float)):
        return f"{v:.4g}"
    return str(v)


# ============================================================
# 3. DISPLAY MODELS — material & section tables (single source)
# ============================================================
def build_material_display():
    """Material properties as (property, value, unit) rows in display units."""
    return [
        ("Material name", MATERIAL["name"], "—"),
        ("Young's modulus  E", to_display_stress(MATERIAL["E"]), UNIT_LABELS["stress"]),
        ("Shear modulus  G", to_display_stress(MATERIAL["G"]), UNIT_LABELS["stress"]),
        ("Yield strength  Fy", to_display_stress(MATERIAL["Fy"]), UNIT_LABELS["stress"]),
        ("Ultimate strength  Fu", to_display_stress(MATERIAL["Fu"]), UNIT_LABELS["stress"]),
        ("Poisson's ratio  nu", MATERIAL["nu"], "—"),
        ("Mass density  rho", MATERIAL["rho"], "kg/m³"),
        ("Thermal coeff.  alpha", MATERIAL["alpha"], "1/°C"),
    ]


def build_section_display():
    """Section properties as (property, value, unit) rows in display units."""
    return [
        ("Section name", section_label(), "—"),
        ("Depth  d", to_display_section_dim(SECTION["d"]), UNIT_LABELS["section_dim"]),
        ("Flange width  bf", to_display_section_dim(SECTION["bf"]), UNIT_LABELS["section_dim"]),
        ("Flange thickness  tf", to_display_section_dim(SECTION["tf"]), UNIT_LABELS["section_dim"]),
        ("Web thickness  tw", to_display_section_dim(SECTION["tw"]), UNIT_LABELS["section_dim"]),
        ("Area  A", to_display_section_area(SECTION["A"]), UNIT_LABELS["area"]),
        ("Strong-axis inertia  Ix", to_display_section_inertia(SECTION["Ix"]), UNIT_LABELS["inertia"]),
        ("Weak-axis inertia  Iy", to_display_section_inertia(SECTION["Iy"]), UNIT_LABELS["inertia"]),
        ("Torsion constant  J", to_display_section_inertia(SECTION["J"]), UNIT_LABELS["inertia"]),
    ]


def material_info_text():
    """Two-part material + section summary for the matplotlib text box."""
    return (
        f"Material: {MATERIAL['name']}\n"
        f"  E = {fmt_value(to_display_stress(MATERIAL['E']))} "
        f"{UNIT_LABELS['stress']},  "
        f"Fy = {fmt_value(to_display_stress(MATERIAL['Fy']))} "
        f"{UNIT_LABELS['stress']}\n"
        f"Section: {section_label()}\n"
        f"  A = {fmt_value(to_display_section_area(SECTION['A']))} "
        f"{UNIT_LABELS['area']},  "
        f"Ix = {fmt_value(to_display_section_inertia(SECTION['Ix']))} "
        f"{UNIT_LABELS['inertia']}"
    )


# ============================================================
# 4. MODEL — members, supports, DOF, local axes, releases
# ============================================================
def classify_member(i, j):
    """Return 'Column' for vertical members, else 'Beam'."""
    xi, yi, zi = nodes[i]
    xj, yj, zj = nodes[j]
    if xi == xj and zi == zj and yi != yj:
        return "Column"
    return "Beam"


members = [(i, j, classify_member(i, j)) for i, j in member_pairs]

# Beta angle: beams = 0°, columns = 90° (RISA/STAAD convention).
BETA_DEG = {
    idx: (90.0 if mtype == "Column" else 0.0)
    for idx, (i, j, mtype) in enumerate(members, start=1)
}


def get_support(node):
    """Support dict: pinned (Ux,Uy,Uz fixed) at base nodes, else free."""
    if node in BASE_NODES:
        return {"Ux": True, "Uy": True, "Uz": True,
                "Rx": False, "Ry": False, "Rz": False}
    return {d: False for d in DOF_NAMES}


supports = {n: get_support(n) for n in nodes}


def build_dof_numbers():
    """Number active DOFs sequentially; restrained DOFs get 'R'."""
    numbers = {}
    counter = 0
    for n in sorted(nodes):
        dof_map = {}
        for d in DOF_NAMES:
            if supports[n][d]:
                dof_map[d] = "R"
            else:
                counter += 1
                dof_map[d] = counter
        numbers[n] = dof_map
    return numbers, counter


node_dof_numbers, ACTIVE_DOF = build_dof_numbers()
TOTAL_DOF = len(nodes) * 6
RESTRAINED_DOF = TOTAL_DOF - ACTIVE_DOF


def local_axes(i, j, beta_deg=0.0):
    """Return (L1, L2, L3) unit vectors for member i-j at beta_deg.

    Reference is always Global Y; for vertical members an internal 90°
    is added to beta so beta=90° yields the conventional physical axes.
    """
    p_i = np.array(nodes[i], dtype=float)
    p_j = np.array(nodes[j], dtype=float)
    vec = p_j - p_i
    length = np.linalg.norm(vec)
    local_1 = vec / length

    global_y = np.array([0.0, 1.0, 0.0])
    global_x = np.array([1.0, 0.0, 0.0])

    if abs(np.dot(local_1, global_y)) > 0.9999:
        ref = global_x
        beta_deg += 90.0            # internal compensation
    else:
        ref = global_y

    local_3 = np.cross(local_1, ref)
    local_3 /= np.linalg.norm(local_3)
    local_2 = np.cross(local_3, local_1)
    local_2 /= np.linalg.norm(local_2)

    beta = np.radians(beta_deg)
    local_2_rot = (local_2 * np.cos(beta)
                   + np.cross(local_1, local_2) * np.sin(beta)
                   + local_1 * np.dot(local_1, local_2) * (1 - np.cos(beta)))
    local_2_rot /= np.linalg.norm(local_2_rot)
    local_3_rot = np.cross(local_1, local_2_rot)
    local_3_rot /= np.linalg.norm(local_3_rot)

    return local_1, local_2_rot, local_3_rot


member_local_axes = {
    idx: local_axes(i, j, BETA_DEG[idx])
    for idx, (i, j, _t) in enumerate(members, start=1)
}


def default_release(mtype):
    """End releases: Beam Mz released at both ends; columns fixed."""
    end = {d: False for d in LOCAL_DOF_NAMES}
    if mtype == "Beam":
        end["Mz"] = True
    return end


member_releases = {
    idx: {"i": default_release(mtype), "j": default_release(mtype)}
    for idx, (i, j, mtype) in enumerate(members, start=1)
}


def total_member_length():
    """Sum of all member lengths (m) from node geometry."""
    return sum(float(np.linalg.norm(
        np.array(nodes[j]) - np.array(nodes[i]))) for i, j in member_pairs)


TOTAL_MEMBER_LENGTH = total_member_length()
# Mass estimate assumes every member uses the base SECTION and MATERIAL.
TOTAL_MASS_KG = MATERIAL["rho"] * SECTION["A"] * TOTAL_MEMBER_LENGTH


# ============================================================
# 5. EXCEL OUTPUT
# ============================================================
wb = openpyxl.Workbook()
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
body_font = Font(name="Arial")
center = Alignment(horizontal="center")


def style_header(ws, headers):
    """Write and style a header row."""
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center


def style_body(ws, ncols):
    """Apply body font/alignment to all data rows."""
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ncols):
        for cell in row:
            cell.font = body_font
            cell.alignment = center


def write_property_sheet(ws, title, rows):
    """Write a property/value/unit table (used by Materials & Sections)."""
    style_header(ws, [title, "Value", "Unit"])
    for name, value, unit in rows:
        ws.append([name, value if isinstance(value, (int, float)) else value, unit])
    style_body(ws, 3)


def build_excel():
    """Create all worksheets and save to EXCEL_PATH."""
    ws_nodes = wb.active
    ws_nodes.title = "Nodes"
    u = UNIT_LABELS["length"]
    style_header(ws_nodes, ["Node", f"X ({u})", f"Y ({u})", f"Z ({u})"])
    for n, (x, y, z) in nodes.items():
        ws_nodes.append([n, to_display_length(x), to_display_length(y), to_display_length(z)])
    style_body(ws_nodes, 4)

    ws_mem = wb.create_sheet("Member Incidences")
    style_header(ws_mem, ["Member", "i (Start Node)", "j (End Node)",
                          "Type", "Beta (deg)", "Material", "Section"])
    for idx, (i, j, mtype) in enumerate(members, start=1):
        ws_mem.append([idx, i, j, mtype, BETA_DEG[idx],
                       MATERIAL["name"], section_label()])
    style_body(ws_mem, 7)

    ws_sup = wb.create_sheet("Supports")
    style_header(ws_sup, ["Node", "Support Type", "Ux", "Uy", "Uz", "Rx", "Ry", "Rz"])
    for n in sorted(nodes):
        stype = "Pinned" if n in BASE_NODES else "Free"
        row = [n, stype] + ["Fixed" if supports[n][d] else "Free" for d in DOF_NAMES]
        ws_sup.append(row)
    style_body(ws_sup, 8)

    ws_dof = wb.create_sheet("Node DOF")
    style_header(ws_dof, ["Node"] + DOF_NAMES)
    for n in sorted(nodes):
        ws_dof.append([n] + [node_dof_numbers[n][d] for d in DOF_NAMES])
    style_body(ws_dof, 7)
    ws_dof.append([])
    ws_dof.append(["Total DOF:", TOTAL_DOF])
    ws_dof.append(["Active DOF:", ACTIVE_DOF])
    ws_dof.append(["Restrained DOF:", RESTRAINED_DOF])

    ws_axes = wb.create_sheet("Member Local Axes")
    style_header(ws_axes, ["Member", "i", "j", "Beta (deg)",
                           "Local-1 X", "Local-1 Y", "Local-1 Z",
                           "Local-2 X", "Local-2 Y", "Local-2 Z",
                           "Local-3 X", "Local-3 Y", "Local-3 Z"])
    for idx, (i, j, mtype) in enumerate(members, start=1):
        l1, l2, l3 = member_local_axes[idx]
        ws_axes.append([idx, i, j, BETA_DEG[idx],
                        round(l1[0], 3), round(l1[1], 3), round(l1[2], 3),
                        round(l2[0], 3), round(l2[1], 3), round(l2[2], 3),
                        round(l3[0], 3), round(l3[1], 3), round(l3[2], 3)])
    style_body(ws_axes, 13)

    ws_rel = wb.create_sheet("Member Releases")
    style_header(ws_rel, ["Member", "End", "Fx", "Fy", "Fz", "Mx", "My", "Mz"])
    for idx, (i, j, mtype) in enumerate(members, start=1):
        for end_label, node_ref in (("i", i), ("j", j)):
            rel = member_releases[idx][end_label]
            row = [idx, f"{end_label} (Node {node_ref})"] + \
                  ["Released" if rel[d] else "Fixed" for d in LOCAL_DOF_NAMES]
            ws_rel.append(row)
    style_body(ws_rel, 8)

    ws_mat = wb.create_sheet("Materials")
    write_property_sheet(ws_mat, "Material Property", build_material_display())

    ws_sec = wb.create_sheet("Sections")
    write_property_sheet(ws_sec, "Section Property", build_section_display())

    ws_sum = wb.create_sheet("Summary")
    style_header(ws_sum, ["Item", "Value"])
    summary_rows = [
        ("Unit system", UNIT_SYSTEM),
        ("Material", MATERIAL["name"]),
        ("Section", section_label()),
        (f"Cube side length ({u})", to_display_length(L)),
        (f"Total member length ({u})", to_display_length(TOTAL_MEMBER_LENGTH)),
        (f"Total mass estimate ({UNIT_LABELS['mass']})",
         round(to_display_mass(TOTAL_MASS_KG), 2)),
        ("Nodes", len(nodes)),
        ("Members", len(members)),
        ("Pinned supports", len(BASE_NODES)),
        ("Total DOF", TOTAL_DOF),
        ("Active DOF", ACTIVE_DOF),
        ("Restrained DOF", RESTRAINED_DOF),
    ]
    for item, value in summary_rows:
        ws_sum.append([item, value])
    style_body(ws_sum, 2)

    wb.save(EXCEL_PATH)
    print(f"Saved Excel: {EXCEL_PATH}")


# ============================================================
# 6. DIAGRAM (matplotlib 3D render)
# ============================================================
def build_cube_figure():
    """Build and return the 3D cube figure (render in display units)."""
    fig = plt.figure(figsize=(11, 10))
    ax = fig.add_subplot(111, projection="3d")

    # members
    for i, j, mtype in members:
        x_vals = [to_plot(nodes[i][0]), to_plot(nodes[j][0])]
        z_vals = [to_plot(nodes[i][2]), to_plot(nodes[j][2])]
        y_vals = [to_plot(nodes[i][1]), to_plot(nodes[j][1])]
        color = "darkorange" if mtype == "Column" else "steelblue"
        ax.plot(x_vals, z_vals, y_vals, color=color, linewidth=2)

    # nodes + DOF labels
    for n, (x, y, z) in nodes.items():
        px, pz, py = to_plot(x), to_plot(z), to_plot(y)
        ax.scatter([px], [pz], [py], color="red", s=60, depthshade=True)
        ax.text(px, pz, py, f"  {n}", color="black",
                fontsize=11, fontweight="bold")
        dof_map = node_dof_numbers[n]
        dof_str = "\n".join([
            f"Ux:{dof_map['Ux']} Uy:{dof_map['Uy']} Uz:{dof_map['Uz']}",
            f"Rx:{dof_map['Rx']} Ry:{dof_map['Ry']} Rz:{dof_map['Rz']}",
        ])
        ax.text(px, pz, py - to_plot(0.35), dof_str,
                color="darkgreen", fontsize=7, fontweight="bold")

    # pinned support triangles
    for n in BASE_NODES:
        x, y, z = nodes[n]
        ax.scatter([to_plot(x)], [to_plot(z)], [to_plot(y)],
                   marker="^", color="black", s=200, zorder=5)

    # local axes triads + beta labels
    for idx, (i, j, mtype) in enumerate(members, start=1):
        mx = to_plot((nodes[i][0] + nodes[j][0]) / 2)
        mz = to_plot((nodes[i][2] + nodes[j][2]) / 2)
        my = to_plot((nodes[i][1] + nodes[j][1]) / 2)
        l1, l2, l3 = member_local_axes[idx]
        s = to_plot(0.6)
        ax.quiver(mx, mz, my, l1[0], l1[2], l1[1], length=s, color="red", linewidth=1.2)
        ax.quiver(mx, mz, my, l2[0], l2[2], l2[1], length=s, color="green", linewidth=1.2)
        ax.quiver(mx, mz, my, l3[0], l3[2], l3[1], length=s, color="blue", linewidth=1.2)
        ax.text(mx + l1[0]*s*1.15, mz + l1[2]*s*1.15, my + l1[1]*s*1.15,
                "L1", color="red", fontsize=8, fontweight="bold")
        ax.text(mx + l2[0]*s*1.15, mz + l2[2]*s*1.15, my + l2[1]*s*1.15,
                "L2", color="green", fontsize=8, fontweight="bold")
        ax.text(mx + l3[0]*s*1.15, mz + l3[2]*s*1.15, my + l3[1]*s*1.15,
                "L3", color="blue", fontsize=8, fontweight="bold")
        ax.text(mx, mz, my + to_plot(0.35), f"β={BETA_DEG[idx]:.0f}°",
                color="purple", fontsize=8, fontweight="bold")

    # global axis triad
    axis_len = to_plot(L * 0.4)
    ax.quiver(0, 0, 0, axis_len, 0, 0, color="black", linewidth=2, arrow_length_ratio=0.12)
    ax.quiver(0, 0, 0, 0, axis_len, 0, color="black", linewidth=2, arrow_length_ratio=0.12)
    ax.quiver(0, 0, 0, 0, 0, axis_len, color="black", linewidth=2, arrow_length_ratio=0.12)
    ax.text(axis_len * 1.05, 0, 0, "Global X", fontsize=9, fontweight="bold")
    ax.text(0, axis_len * 1.05, 0, "Global Z", fontsize=9, fontweight="bold")
    ax.text(0, 0, axis_len * 1.05, "Global Y", fontsize=9, fontweight="bold")

    u = UNIT_LABELS["length"]
    ax.set_xlabel(f"X ({u}) - Lateral", fontsize=11, labelpad=10)
    ax.set_ylabel(f"Z ({u}) - Lateral", fontsize=11, labelpad=10)
    ax.set_zlabel(f"Y ({u}) - Vertical", fontsize=11, labelpad=10)

    ax.set_title(
        f"6x6x6 Cube Model (Rev. 2)  |  Material: {MATERIAL['name']}  |  "
        f"Section: {section_label()}  |  Units: {UNIT_SYSTEM}\n"
        "β (beams=0°, columns=90°)  |  L1/L2/L3 = Local axes  |  ▲ = Pinned",
        fontsize=10, fontweight="bold")

    lim = to_plot(L)
    ax.set_xlim(-to_plot(1.5), lim + to_plot(1.5))
    ax.set_ylim(-to_plot(1.5), lim + to_plot(1.5))
    ax.set_zlim(-to_plot(1.0), lim + to_plot(1.0))
    ax.set_box_aspect([1, 1, 1])
    ax.view_init(elev=20, azim=-60)

    legend_elems = [
        Line2D([0], [0], color="steelblue", lw=2, label="Beam (β=0°)"),
        Line2D([0], [0], color="darkorange", lw=2, label="Column (β=90°)"),
        Line2D([0], [0], marker="^", color="black", lw=0, markersize=10, label="Pinned Support"),
        Line2D([0], [0], color="red", lw=1, label="Local-1 (axial)"),
        Line2D([0], [0], color="green", lw=1, label="Local-2"),
        Line2D([0], [0], color="blue", lw=1, label="Local-3"),
    ]
    ax.legend(handles=legend_elems, loc="upper left", fontsize=8)

    # material + section text box
    ax.text2D(0.015, -0.02, material_info_text(), transform=ax.transAxes,
              fontsize=8, family="monospace", va="top", ha="left",
              bbox=dict(boxstyle="round,pad=0.4", facecolor="white",
                        edgecolor="0.6", alpha=0.9))

    fig.tight_layout()
    return fig


fig_embed = build_cube_figure()
buf = io.BytesIO()
fig_embed.savefig(buf, format="png", dpi=150, bbox_inches="tight")
plt.close(fig_embed)
buf.seek(0)
IMG_B64 = base64.b64encode(buf.read()).decode("ascii")


# ============================================================
# 7. HTML OUTPUT
# ============================================================
def property_table_html(title, rows):
    """Render a (property, value, unit) list as an HTML table."""
    body = "".join(
        f"<tr><td>{name}</td><td>{fmt_value(value)}</td><td>{unit}</td></tr>"
        for name, value, unit in rows
    )
    return (f'<section class="panel"><h2>{title}</h2>'
            f'<table><tr><th>Property</th><th>Value</th><th>Unit</th></tr>'
            f"{body}</table></section>")


generated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
u = UNIT_LABELS["length"]

node_rows = "".join(
    f"<tr><td>{n}</td><td>{to_display_length(x):.1f}</td>"
    f"<td>{to_display_length(y):.1f}</td><td>{to_display_length(z):.1f}</td></tr>"
    for n, (x, y, z) in nodes.items()
)

member_rows = ""
for idx, (i, j, mtype) in enumerate(members, start=1):
    cls = "col-beam" if mtype == "Beam" else "col-column"
    member_rows += (f"<tr><td>{idx}</td><td>{i}</td><td>{j}</td>"
                    f"<td class='{cls}'>{mtype}</td><td>{BETA_DEG[idx]:.1f}</td>"
                    f"<td>{MATERIAL['name']}</td><td>{section_label()}</td></tr>")

support_rows = ""
for n in sorted(nodes):
    stype = "Pinned" if n in BASE_NODES else "Free"
    cells = "".join(f"<td>{'Fixed' if supports[n][d] else 'Free'}</td>" for d in DOF_NAMES)
    support_rows += f"<tr><td>{n}</td><td>{stype}</td>{cells}</tr>"

dof_rows = ""
for n in sorted(nodes):
    vals = "".join(f"<td>{node_dof_numbers[n][d]}</td>" for d in DOF_NAMES)
    dof_rows += f"<tr><td>{n}</td>{vals}</tr>"

material_section_html = (
    property_table_html("Material — " + MATERIAL["name"], build_material_display())
    + property_table_html("Section — " + section_label(), build_section_display())
)

CSS = """
:root { --ink:#0b2531; --muted:#5b7180; --line:#d7e3e6;
        --teal:#0e7c86; --teal-dark:#0b4f6c; }
* { box-sizing: border-box; }
body { margin:0; font-family:'Segoe UI', Arial, sans-serif;
       color:var(--ink); background:#f2fafa; line-height:1.55; }
header.hero { background:linear-gradient(135deg,#0e7c86 0%, #0b4f6c 100%);
              color:#fff; padding:32px 28px 60px 28px; position:relative; }
header.hero .kicker { display:inline-block;
       background:rgba(255,255,255,0.15); padding:4px 10px;
       border-radius:20px; font-size:11px; letter-spacing:0.08em;
       text-transform:uppercase; }
header.hero h1 { margin:12px 0 6px 0; font-size:32px; }
header.hero .sub { opacity:0.85; font-size:14px; }
main { max-width:1200px; margin:0 auto; padding:24px 20px 40px; }
section.panel { background:#fff; border-radius:16px; padding:22px 24px;
                margin-bottom:22px; box-shadow:0 6px 24px rgba(11,79,108,0.08); }
h2 { margin:0 0 12px 0; color:var(--teal-dark); font-size:20px; }
.chart { max-width:100%; height:auto; border-radius:10px;
         display:block; margin:0 auto; }
table { border-collapse:collapse; width:100%; font-size:13px; margin:10px 0; }
th, td { border-bottom:1px solid var(--line); padding:9px 10px; text-align:left; }
th { background:#effaf9; color:var(--teal-dark);
     text-transform:uppercase; font-size:10.5px; letter-spacing:0.05em; }
tr:last-child td { border-bottom:none; }
.col-beam   { color:#2a7de1; font-weight:700; }
.col-column { color:#c25e00; font-weight:700; }
footer { text-align:center; color:var(--muted); font-size:11.5px; padding:14px; }
"""

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>6x6x6 Cube Model &mdash; Rev. 2</title>
<style>
{css}
</style>
</head>
<body>
<header class="hero">
  <span class="kicker">Rev. 2 &middot; {unit_system}</span>
  <h1>6 &times; 6 &times; 6 Cube Frame</h1>
  <div class="sub">Material: {material} &middot; Section: {section} &middot;
    Supports &middot; Beta (beams 0&deg;, columns 90&deg;) &middot; Local axes &middot;
    DOF &middot; Releases &middot; generated {generated}</div>
</header>
<main>
  <section class="panel">
    <h2>3D Diagram</h2>
    <img class="chart" src="data:image/png;base64,{img_b64}" alt="cube Rev2">
  </section>
{material_section}
  <section class="panel"><h2>Nodes</h2>
    <table><tr><th>Node</th><th>X ({u})</th><th>Y ({u})</th><th>Z ({u})</th></tr>
    {node_rows}</table></section>
  <section class="panel"><h2>Member Incidences</h2>
    <table><tr><th>Member</th><th>i</th><th>j</th><th>Type</th><th>&beta; (&deg;)</th>
    <th>Material</th><th>Section</th></tr>
    {member_rows}</table></section>
  <section class="panel"><h2>Supports</h2>
    <table><tr><th>Node</th><th>Type</th><th>Ux</th><th>Uy</th><th>Uz</th>
    <th>Rx</th><th>Ry</th><th>Rz</th></tr>{support_rows}</table></section>
  <section class="panel"><h2>Node DOF</h2>
    <table><tr><th>Node</th><th>Ux</th><th>Uy</th><th>Uz</th>
    <th>Rx</th><th>Ry</th><th>Rz</th></tr>{dof_rows}</table></section>
</main>
<footer>Generated by <code>Rev2.py</code></footer>
</body>
</html>
"""

html = HTML_TEMPLATE.format(
    css=CSS, unit_system=UNIT_SYSTEM, material=MATERIAL["name"],
    section=section_label(), generated=generated, img_b64=IMG_B64,
    material_section=material_section_html,
    node_rows=node_rows, member_rows=member_rows,
    support_rows=support_rows, dof_rows=dof_rows, u=u,
)

with open(HTML_PATH, "w", encoding="utf-8") as f:
    f.write(html)
print(f"Saved HTML:  {HTML_PATH}")


# ============================================================
# 8. RUN + POP-UP
# ============================================================
build_excel()

fig_show = build_cube_figure()
print("Opening cube plot window — close it to exit.")
plt.show()


# ============================================================
# NOTES
# ============================================================
# WHAT CHANGED vs Rev1:
#   1. Unit system: UNIT_SYSTEM constant ("metric"/"imperial") with a single
#      UNITS dict of NIST SP 811 factors; all display goes through helpers
#      (to_display_length/force/stress + section/mass + inverses). Physics
#      stays in SI internally; only display/input converts.
#   2. Material: A36 Steel declared once (E, Fy, Fu, rho, nu, G, alpha) from
#      the RISA Materials Library (Metric), applied to all 12 members.
#   3. Section: W 100x100 metric (== W 4x13) from the AISC DB v16.0, applied
#      to all members; SECTION_OVERRIDES allows per-member override.
#   4. New outputs: Excel gains "Materials" + "Sections" sheets; "Member
#      Incidences" gains Material/Section columns; "Summary" states unit
#      system, material, section and a total mass estimate.
#   5. HTML gains material + section tables and shows the unit system in the
#      header; the matplotlib title + text box show material + section names.
#   6. Section label switches: "W 100x100" (metric) / "W 4x13" (imperial).
#
# SWITCH UNITS: set UNIT_SYSTEM = "imperial" (or "metric") at the top and rerun.
#
# ASSUMPTIONS:
#   - Section properties source: AISC Shapes Database v16.0 (row W4X13 /
#     W100X19.3), cross-checked with the MEsteel metric W 100x100 x 19 table.
#   - Material source: RISA Materials Library (Metric), "Hot Rolled" A36 Gr.36.
#   - Mass estimate = rho * A * total member length, assuming every member uses
#     the base SECTION and MATERIAL (SECTION_OVERRIDES empty by default).
#   - Only display labels/values switch with UNIT_SYSTEM; internal SI state,
#     geometry, DOF numbering, local axes and releases are unchanged.








